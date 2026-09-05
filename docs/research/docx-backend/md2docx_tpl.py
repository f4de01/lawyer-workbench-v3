"""Chain 1: Markdown(全文) -> DOCX, 以官方模板为版式载体 (python-docx, 离线).
约定: '# ' 标题; 普通行=正文段; ':right: ' 前缀=右对齐; 管道表=表格; 空行=段落分隔.
"""
import sys, re, copy, datetime, argparse
import docx
from docx.shared import Pt, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

def set_font(run, name, size_pt):
    run.font.name = name; run.font.size = Pt(size_pt)
    rpr = run._r.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {}); rpr.insert(0, rf)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a), name)

def para(doc, text, font='仿宋', size=16, align=None, indent=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = Pt(25); pf.line_spacing_rule = 4  # exact 500 twips
    if indent and align is None: pf.first_line_indent = Twips(640)
    if align: p.alignment = align
    if text:
        set_font(p.add_run(text), font, size)
    return p

def convert(md_text, template, out, strip_meta=True, keep_trailing_para=True, extra_blank=0):
    doc = docx.Document(template)
    body = doc.element.body
    for child in list(body):
        if not child.tag.endswith('sectPr'): body.remove(child)
    lines = md_text.splitlines(); i = 0; last_was_table = False
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip(): i += 1; continue
        if ln.startswith('# '):
            para(doc, ln[2:].strip(), '方正小标宋简体', 18, WD_ALIGN_PARAGRAPH.CENTER, indent=False); last_was_table=False
        elif ln.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', c) for c in cells): rows.append(cells)
                i += 1
            t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows)); t.style = 'Table Grid'
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    cell = t.cell(r, c); cell.paragraphs[0].text = ''
                    if val: set_font(cell.paragraphs[0].add_run(val), '仿宋', 16)
            last_was_table = True; continue
        elif ln.startswith(':right: '):
            para(doc, ln[8:].strip(), align=WD_ALIGN_PARAGRAPH.RIGHT); last_was_table=False
        else:
            para(doc, ln.strip()); last_was_table=False
        i += 1
    for _ in range(extra_blank): para(doc, '')
    if last_was_table and keep_trailing_para: para(doc, '')
    if not keep_trailing_para:  # 事故5: 表格直接接 sectPr
        for p in reversed(list(body)):
            if p.tag.endswith('}p') and not ''.join(p.itertext()).strip(): body.remove(p)
            else: break
    if strip_meta:
        cp = doc.core_properties; now = datetime.datetime.now()
        cp.author = ''; cp.last_modified_by = ''; cp.revision = 1
        lp = cp._element.find('{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}lastPrinted')
        if lp is not None: cp._element.remove(lp)  # python-docx 不接受 None，须直接删元素
        cp.created = now; cp.modified = now
    doc.save(out)

if __name__ == '__main__':
    ap = argparse.ArgumentParser(); ap.add_argument('md'); ap.add_argument('template'); ap.add_argument('out')
    ap.add_argument('--keep-meta', action='store_true'); ap.add_argument('--no-trailing', action='store_true'); ap.add_argument('--extra-blank', type=int, default=0)
    a = ap.parse_args()
    convert(open(a.md, encoding='utf-8').read(), a.template, a.out, not a.keep_meta, not a.no_trailing, a.extra_blank)
    print('wrote', a.out)
